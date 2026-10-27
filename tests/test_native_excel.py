"""Native Excel compatibility against frozen legacy workbooks and payloads."""

from __future__ import annotations

import builtins
import datetime as dt
import json
from pathlib import Path
import subprocess
import sys
from typing import Any

import openpyxl
import pandas as pd
import polars as pl
import pytest
from xlsxwriter.exceptions import InvalidWorksheetName

from voiage import _tabular_io
from voiage.ecosystem_integration import DataFormatConnector
from voiage.health_economics import HealthEconomicsAnalysis, HealthState, Treatment

ROOT = Path(__file__).parents[1]
FIXTURES = ROOT / "tests/fixtures/tabular_io/excel"
BASELINE = json.loads(
    (
        ROOT / "conductor/tracks/pandas_free_io_20261026/phase-3-baseline.json"
    ).read_text()
)


def _encode(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            "mapping": [[_encode(key), _encode(item)] for key, item in value.items()]
        }
    if isinstance(value, (list, tuple)):
        return [_encode(item) for item in value]
    if pd.isna(value):
        return {"missing": True}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"datetime": value.isoformat()}
    return {"type": type(value).__name__, "value": value}


@pytest.mark.parametrize("filename", sorted(BASELINE["fixtures"]))
def test_native_reader_retains_frozen_payload(filename: str) -> None:
    expected = BASELINE["fixtures"][filename]
    if expected["warnings"]:
        with pytest.warns(UserWarning, match="Error importing health data"):
            result = DataFormatConnector().import_health_data(
                str(FIXTURES / filename), "generic"
            )
    else:
        result = DataFormatConnector().import_health_data(
            str(FIXTURES / filename), "generic"
        )
    assert _encode(result) == expected["payload"]


def _analysis() -> HealthEconomicsAnalysis:
    analysis = HealthEconomicsAnalysis(willingness_to_pay=50_000.0, currency="AUD")
    analysis.add_treatment(
        Treatment(
            name="Standard Care",
            description="Baseline",
            effectiveness=0.75,
            cost_per_cycle=100.0,
            cycles_required=10,
        )
    )
    analysis.add_treatment(
        Treatment(
            name="Drug A",
            description="Intervention",
            effectiveness=0.9,
            cost_per_cycle=150.0,
            cycles_required=10,
        )
    )
    analysis.add_health_state(
        HealthState(
            state_id="stable",
            description="Stable disease",
            utility=0.85,
            cost=50.0,
            duration=1.0,
        )
    )
    return analysis


def test_native_writer_retains_legacy_sheets_and_cells(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    DataFormatConnector().export_health_data(_analysis(), str(path), "xlsx")
    actual = openpyxl.load_workbook(path, data_only=True)
    expected = openpyxl.load_workbook(FIXTURES / "legacy-export.xlsx", data_only=True)
    assert actual.sheetnames == expected.sheetnames
    for name in expected.sheetnames:
        assert list(actual[name].values) == list(expected[name].values)
        assert [cell.number_format for cell in actual[name][2]] == [
            cell.number_format for cell in expected[name][2]
        ]
        assert [cell.font.bold for cell in actual[name][1]] == [
            cell.font.bold for cell in expected[name][1]
        ]


def test_native_writer_retains_formula_cache_urls_and_non_finites(
    tmp_path: Path,
) -> None:
    path = tmp_path / "special.xlsx"
    frame = pl.DataFrame(
        {
            "text": ["=1+1", "https://example.test", "literal", "last"],
            "number": [1.25, float("inf"), float("-inf"), float("nan")],
            "flag": [True, False, True, False],
        }
    )
    _tabular_io.write_excel_tables({"values": frame}, path)
    for cached in [True, False]:
        actual = openpyxl.load_workbook(path, data_only=cached)["values"]
        expected = openpyxl.load_workbook(
            FIXTURES / "writer-special.xlsx", data_only=cached
        )["values"]
        assert list(actual.values) == list(expected.values)
        assert actual["A3"].hyperlink is None
        assert [[c.number_format for c in row] for row in actual] == [
            [c.number_format for c in row] for row in expected
        ]


@pytest.mark.parametrize(
    "frame", [pl.DataFrame(), pl.DataFrame(schema={"value": pl.Float64})]
)
def test_native_writer_empty_tables(tmp_path: Path, frame: pl.DataFrame) -> None:
    path = tmp_path / "empty.xlsx"
    _tabular_io.write_excel_tables({"empty": frame}, path)
    sheet = openpyxl.load_workbook(path)["empty"]
    assert list(sheet.values) == ([tuple(frame.columns)] if frame.width else [])


@pytest.mark.parametrize(
    ("module", "operation"), [("python_calamine", "read"), ("xlsxwriter", "write")]
)
def test_missing_native_engine_fails_soft_without_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, module: str, operation: str
) -> None:
    original_import = builtins.__import__

    def blocked(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == module:
            raise ModuleNotFoundError(name)
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked)
    connector = DataFormatConnector()
    if operation == "read":
        with pytest.warns(UserWarning, match=r"requires.*voiage\[excel\]"):
            result = connector.import_health_data(str(FIXTURES / "mixed.xlsx"))
        assert result == {}
    else:
        with pytest.warns(UserWarning, match=r"requires.*voiage\[excel\]"):
            result = connector.export_health_data(
                _analysis(), str(tmp_path / "missing.xlsx"), "xlsx"
            )
        assert result is None
        assert not (tmp_path / "missing.xlsx").exists()


def test_native_excel_without_pandas_or_openpyxl(tmp_path: Path) -> None:
    script = r"""
import importlib.abc
from pathlib import Path
import sys
class BlockLegacy(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'pandas', 'openpyxl', 'fastexcel'}:
            raise ModuleNotFoundError(fullname)
sys.meta_path.insert(0, BlockLegacy())
from voiage.ecosystem_integration import DataFormatConnector
from voiage.health_economics import HealthEconomicsAnalysis, Treatment
connector = DataFormatConnector()
loaded = connector.import_health_data(sys.argv[1])
assert loaded['shape'] == (4, 6), loaded
assert loaded['data'][0]['mixed'] == 1
analysis = HealthEconomicsAnalysis()
analysis.add_treatment(Treatment(name='Drug', description='Native round-trip', effectiveness=.8, cost_per_cycle=10., cycles_required=1))
connector.export_health_data(analysis, sys.argv[2], 'xlsx')
assert Path(sys.argv[2]).exists()
again = connector.import_health_data(sys.argv[2])
assert again['data'][0]['name'] == 'Drug', again
assert not any(name.split('.')[0] in {'pandas','openpyxl','fastexcel'} for name in sys.modules)
"""
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(FIXTURES / "mixed.xlsx"),
            str(tmp_path / "native.xlsx"),
        ],
        capture_output=True,
        text=True,
        timeout=90,
        check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        ([True, None, False], [1.0, float("nan"), 0.0]),
        (["true", "FALSE"], [True, False]),
        ([dt.datetime(2026, 1, 1), "text"], [dt.datetime(2026, 1, 1), "text"]),
        ([True, "text"], [True, "text"]),
    ],
)
def test_native_reader_full_column_scalar_policy(
    tmp_path: Path, values: list[Any], expected: list[Any]
) -> None:
    path = tmp_path / "cells.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["value", "retain_row"])
    for value in values:
        sheet.append([value, "present"])
    workbook.save(path)
    workbook.close()
    loaded = DataFormatConnector().import_health_data(str(path), "generic")
    assert _encode([r["value"] for r in loaded["data"]]) == _encode(expected)


def test_native_reader_header_only_and_invalid_workbook(tmp_path: Path) -> None:
    path = tmp_path / "header.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["value"])
    workbook.save(path)
    workbook.close()
    frame, headers = _tabular_io.read_excel(path)
    assert headers == ["value"]
    assert frame.shape == (0, 1)
    path.write_text("invalid workbook")
    with pytest.warns(UserWarning, match="Error importing health data"):
        result = DataFormatConnector().import_health_data(str(path))
    assert result == {}
    # A failed read must release its native handle as well as a successful read.
    path.unlink()
    assert not path.exists()


def test_native_writer_rejects_invalid_extension_before_creation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "output.csv"
    with pytest.raises(ValueError, match="Invalid extension"):
        _tabular_io.write_excel_tables({"data": pl.DataFrame({"a": [1]})}, path)
    assert not path.exists()


def test_native_writer_closes_workbook_on_failure(tmp_path: Path) -> None:
    path = tmp_path / "invalid.xlsx"
    with pytest.raises(InvalidWorksheetName, match="Invalid Excel character"):
        _tabular_io.write_excel_tables({"invalid/name": pl.DataFrame({"a": [1]})}, path)
    path.unlink()
    assert not path.exists()


@pytest.mark.parametrize("whitespace", [" ", "\t", "\n", " padded ", " 001 "])
@pytest.mark.parametrize("surrounding", [["a", "b"], [1, 2], [True, False]])
def test_native_reader_preserves_whitespace_only_rows(
    tmp_path: Path, whitespace: str, surrounding: list[Any]
) -> None:
    path = tmp_path / "whitespace.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["value"])
    for value in [surrounding[0], whitespace, surrounding[1]]:
        workbook.active.append([value])
    workbook.save(path)
    workbook.close()
    expected = pd.read_excel(path, engine="openpyxl")
    actual = DataFormatConnector().import_health_data(str(path), "generic")
    assert actual["shape"] == expected.shape
    assert _encode(actual["data"]) == _encode(expected.to_dict("records"))


@pytest.mark.parametrize("invalid_xml", [False, True])
def test_native_reader_private_copy_cleanup_and_source_immutability(
    tmp_path: Path, invalid_xml: bool, monkeypatch: pytest.MonkeyPatch
) -> None:
    import tempfile
    from zipfile import ZipFile

    path = tmp_path / "source.xlsx"
    with ZipFile(FIXTURES / "mixed.xlsx") as original, ZipFile(path, "w") as target:
        for info in original.infolist():
            content = original.read(info)
            if invalid_xml and info.filename == "xl/worksheets/sheet1.xml":
                content = b'<!DOCTYPE worksheet [<!ENTITY unsafe "x">]>' + content
            target.writestr(info, content)
    before = path.read_bytes()
    created = []
    original_temporary_directory = tempfile.TemporaryDirectory

    def owned_directory(*args: Any, **kwargs: Any) -> Any:
        directory = original_temporary_directory(*args, dir=tmp_path, **kwargs)
        created.append(Path(directory.name))
        return directory

    monkeypatch.setattr(_tabular_io.tempfile, "TemporaryDirectory", owned_directory)
    if invalid_xml:
        with pytest.warns(UserWarning, match="DTDForbidden"):
            result = DataFormatConnector().import_health_data(str(path), "generic")
        assert result == {}
    else:
        result = DataFormatConnector().import_health_data(str(path), "generic")
        assert result["shape"] == (4, 6)
    assert created
    assert all(not directory.exists() for directory in created)
    assert path.read_bytes() == before
